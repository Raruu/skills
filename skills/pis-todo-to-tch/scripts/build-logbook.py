#!/usr/bin/env python3
"""
build-logbook.py — Build a Polinema internship log book (DOCX + PDF)
from Profile Plus todos and attendance records.

Period resolution (--input):
    ""                               current month, 1st -> today
    "September" / "September 2026"   that month (past: full; current: 1st -> today)
    "September/minggu 2"             calendar week 2 (Mon-Sun), clipped to month
    "September/minggu 2->3"          weeks 2 through 3
    "Tanggal 2026-09-01 -> 2026-09-15"  explicit date range

Data (--data) is a JSON file written by the agent from MCP results:
    {"todos": [{"date": "yyyy-MM-dd", "task": "...", "status": "DONE", "created_at": "..."}],
     "attendance": [{"date": "yyyy-MM-dd", "check_in": "...", "check_out": "...", "status": "HADIR"}]}

Config: personal data is read from `.pis-todo-to-tch.json` in the working
directory (--config overrides the path). If it does not exist yet, it is
created from the shipped example and the script exits 2 so the user can edit
it before generating.

Exit codes: 0 ok | 1 error | 2 config was just created from the example (edit it first)
"""

from __future__ import annotations

import argparse
import calendar
import json
import re
import sys
from copy import deepcopy
from datetime import date, datetime, timedelta
from pathlib import Path

SKILL_DIR = Path(__file__).resolve().parent.parent
ASSETS_DIR = SKILL_DIR / "assets"
DEFAULT_TEMPLATE = ASSETS_DIR / "template.docx"
DEFAULT_LOGO = ASSETS_DIR / "logo.jpeg"
EXAMPLE_CONFIG = SKILL_DIR / "config.example.json"
CONFIG_FILENAME = ".pis-todo-to-tch.json"

HARI = ["Senin", "Selasa", "Rabu", "Kamis", "Jumat", "Sabtu", "Minggu"]
BULAN = ["Januari", "Februari", "Maret", "April", "Mei", "Juni",
         "Juli", "Agustus", "September", "Oktober", "November", "Desember"]

DURASI_RE = re.compile(
    r"\s*:\s*\d+(?:\s*-\s*\d+(?:[.,]\d+)?|[.,]\d+)?\s*(?:jam|menit)\s*$",
    re.IGNORECASE,
)
STATUS_LABEL = {"CUTI": "Cuti", "IZIN": "Izin", "SAKIT": "Sakit"}

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"


# --------------------------------------------------------------------------- #
# period resolution
# --------------------------------------------------------------------------- #

def parse_iso(raw: str) -> date:
    return datetime.strptime(raw.strip(), "%Y-%m-%d").date()


def parse_month(raw: str, today: date) -> tuple[int, int]:
    raw = raw.strip()
    m = re.match(r"^([A-Za-z]+)(?:\s+(\d{4}))?$", raw)
    if not m:
        raise ValueError(f"bulan tidak dikenali: {raw!r}")
    name = m.group(1).capitalize()
    if name not in BULAN:
        raise ValueError(f"bulan tidak dikenali: {raw!r} (contoh: September atau September 2026)")
    month = BULAN.index(name) + 1
    year = int(m.group(2)) if m.group(2) else today.year
    if (year, month) > (today.year, today.month):
        raise ValueError(f"{BULAN[month - 1]} {year} belum berjalan")
    return year, month


def month_bounds(year: int, month: int) -> tuple[date, date]:
    first = date(year, month, 1)
    last = date(year, month, calendar.monthrange(year, month)[1])
    return first, last


def week_bounds(year: int, month: int, w1: int, w2: int) -> tuple[date, date]:
    """Calendar weeks (Mon-Sun); week 1 = the week containing the 1st, clipped to the month."""
    first, last = month_bounds(year, month)
    wk_start = first - timedelta(days=first.weekday())  # Monday of the week holding day 1
    if w1 < 1 or w2 < w1 or w2 > 6:
        raise ValueError("nomor minggu harus 1..6 dan w1 <= w2")
    start = max(wk_start + timedelta(weeks=w1 - 1), first)
    end = min(wk_start + timedelta(weeks=w2) - timedelta(days=1), last)
    if start > last or end < first:
        raise ValueError(f"minggu {w1} di luar {BULAN[month - 1]} {year}")
    return start, end


def parse_period(raw: str, today: date) -> dict:
    raw = (raw or "").strip()
    warnings: list[str] = []

    if not raw:
        y, m = today.year, today.month
        return {
            "mode": "bulan", "start": date(y, m, 1), "end": today,
            "month": (y, m), "weeks": None, "warnings": warnings,
        }

    m = re.match(r"^tanggal\s+(.+?)\s*->\s*(.+)$", raw, re.IGNORECASE)
    if m:
        try:
            start, end = parse_iso(m.group(1)), parse_iso(m.group(2))
        except ValueError:
            raise ValueError("format tanggal harus yyyy-MM-dd, contoh: Tanggal 2026-09-01 -> 2026-09-15")
        if start > end:
            raise ValueError("tanggal awal melebihi tanggal akhir")
        if start > today:
            raise ValueError(f"periode {start} belum berjalan")
        if end > today:
            warnings.append(f"tanggal akhir {end} di masa depan — dipotong ke {today}")
            end = today
        return {
            "mode": "tanggal", "start": start, "end": end,
            "month": (start.year, start.month), "weeks": None, "warnings": warnings,
        }

    m = re.match(r"^(.*?)\s*/\s*minggu\s+(\d+)\s*(?:->\s*(\d+))?$", raw, re.IGNORECASE)
    if m:
        month_part = m.group(1).strip()
        w1 = int(m.group(2))
        w2 = int(m.group(3)) if m.group(3) else w1
        y, mo = parse_month(month_part, today)
        start, end = week_bounds(y, mo, w1, w2)
        if start > today:
            raise ValueError(f"minggu {w1} {BULAN[mo - 1]} {y} belum berjalan")
        if end > today:
            end = today
        return {
            "mode": "minggu", "start": start, "end": end,
            "month": (y, mo), "weeks": (w1, w2), "warnings": warnings,
        }

    y, mo = parse_month(raw, today)
    first, last = month_bounds(y, mo)
    end = last
    if (y, mo) == (today.year, today.month):
        end = today
    return {
        "mode": "bulan", "start": first, "end": end,
        "month": (y, mo), "weeks": None, "warnings": warnings,
    }


def period_labels(period: dict) -> tuple[str, str]:
    """(folder_label, file_label)."""
    start, end = period["start"], period["end"]
    y, mo = period["month"]

    if period["mode"] == "minggu":
        w1, w2 = period["weeks"]
        weeks = f"Minggu {w1}" + (f"-{w2}" if w2 != w1 else "")
        return f"{BULAN[mo - 1]} {y}", f"{BULAN[mo - 1]} {y} - {weeks}"

    if period["mode"] == "tanggal":
        if start.year == end.year and start.month == end.month:
            return (f"{BULAN[start.month - 1]} {start.year}",
                    f"{BULAN[start.month - 1]} {start.year} - {start.day}-{end.day}")
        if start.year == end.year:
            return (f"{BULAN[start.month - 1]} - {BULAN[end.month - 1]} {start.year}",
                    f"{BULAN[start.month - 1]} - {BULAN[end.month - 1]} {start.year} - {start.day}-{end.day}")
        return (f"{BULAN[start.month - 1]} {start.year} - {BULAN[end.month - 1]} {end.year}",
                f"{BULAN[start.month - 1]} {start.year} - {BULAN[end.month - 1]} {end.year} - {start.day}-{end.day}")

    return f"{BULAN[mo - 1]} {y}", f"{BULAN[mo - 1]} {y}"


# --------------------------------------------------------------------------- #
# data assembly
# --------------------------------------------------------------------------- #

def fmt_tanggal(d: date) -> str:
    return f"{HARI[d.weekday()]}, {d.day} {BULAN[d.month - 1]} {d.year}"


def fmt_jam(raw: str | None) -> str:
    if not raw:
        return "-"
    m = re.search(r"(\d{2}):(\d{2})", raw)
    return f"{m.group(1)}:{m.group(2)}" if m else "-"


def strip_durasi(task: str) -> str:
    stripped = DURASI_RE.sub("", task).strip()
    return stripped or task.strip()


def build_days(period: dict, todos: list[dict], attendance: list[dict]) -> list[dict]:
    todos_by_day: dict[str, list[dict]] = {}
    for t in todos:
        todos_by_day.setdefault(str(t.get("date", "")), []).append(t)

    att_by_day: dict[str, dict] = {}
    for a in attendance:
        att_by_day[str(a.get("date", ""))] = a

    days: list[dict] = []
    d = period["start"]
    while d <= period["end"]:
        key = d.isoformat()
        day_todos = sorted(todos_by_day.get(key, []), key=lambda t: str(t.get("created_at") or ""))
        att = att_by_day.get(key)

        items: list[str] = []
        for t in day_todos:
            task = strip_durasi(str(t.get("task", "")))
            if task:
                items.append(task)
        if not items and att and str(att.get("status", "")).upper() in STATUS_LABEL:
            items.append(STATUS_LABEL[str(att["status"]).upper()])

        if d.weekday() < 5 or items or att:
            days.append({
                "date": d,
                "label": fmt_tanggal(d),
                "in": fmt_jam(att.get("check_in")) if att else "-",
                "out": fmt_jam(att.get("check_out")) if att else "-",
                "items": items,
                "no_hours": (not att) or not att.get("check_in") or not att.get("check_out"),
            })
        d += timedelta(days=1)
    return days


# --------------------------------------------------------------------------- #
# DOCX
# --------------------------------------------------------------------------- #

def _replace_in_paragraph(p, target: str, repl: str) -> bool:
    if target not in p.text:
        return False
    text = p.text
    idx = text.index(target)
    end = idx + len(target)
    pos = 0
    first = last = None
    for r in p.runs:
        s, e = pos, pos + len(r.text)
        if s < end and e > idx:
            if first is None:
                first = (r, s)
            last = (r, s, e)
        pos = e
    if first is None:
        return False
    fr, fs = first
    lr, ls, le = last
    before = fr.text[: idx - fs]
    after = lr.text[end - ls:]
    if fr is lr:
        fr.text = before + repl + after
    else:
        fr.text = before + repl
        lr.text = after
        seen = False
        for r in p.runs:
            if r is fr:
                seen = True
                continue
            if r is lr:
                break
            if seen:
                r.text = ""
    return True


def _all_paragraphs(doc):
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for p in c.paragraphs:
                    yield p
    for s in doc.sections:
        for hf in (s.header, s.first_page_header, s.even_page_header,
                   s.footer, s.first_page_footer, s.even_page_footer):
            for p in hf.paragraphs:
                yield p
            for t in hf.tables:
                for row in t.rows:
                    for c in row.cells:
                        for p in c.paragraphs:
                            yield p


def _set_cell_text(tc, text: str) -> None:
    p = tc.find(f"{W_NS}p")
    for r in p.findall(f"{W_NS}r"):
        p.remove(r)
    rpr_src = p.find(f"{W_NS}pPr/{W_NS}rPr")
    r = p.makeelement(f"{W_NS}r", {})
    if rpr_src is not None:
        r.append(deepcopy(rpr_src))
    t = r.makeelement(f"{W_NS}t", {})
    t.set(XML_SPACE, "preserve")
    t.text = text
    r.append(t)
    p.append(r)


def _set_left_alignment(p) -> None:
    """Force left alignment on a paragraph, inserting <w:jc> in schema order.

    The template prototype for the Kegiatan cell is center-aligned; cloned
    paragraphs must be overridden so activity bullets read as a left list.
    """
    ppr = p.find(f"{W_NS}pPr")
    if ppr is None:
        ppr = p.makeelement(f"{W_NS}pPr", {})
        p.insert(0, ppr)
    jc = ppr.find(f"{W_NS}jc")
    if jc is None:
        jc = ppr.makeelement(f"{W_NS}jc", {})
        rpr = ppr.find(f"{W_NS}rPr")
        if rpr is not None:
            rpr.addprevious(jc)
        else:
            ppr.append(jc)
    jc.set(f"{W_NS}val", "left")


def _kegiatan_paragraph(proto_p, text: str):
    p = deepcopy(proto_p)
    _set_left_alignment(p)
    for r in p.findall(f"{W_NS}r"):
        p.remove(r)
    rpr_src = p.find(f"{W_NS}pPr/{W_NS}rPr")
    r = p.makeelement(f"{W_NS}r", {})
    if rpr_src is not None:
        r.append(deepcopy(rpr_src))
    t = r.makeelement(f"{W_NS}t", {})
    t.set(XML_SPACE, "preserve")
    t.text = text
    r.append(t)
    p.append(r)
    return p


def fill_docx(template: Path, out_path: Path, cfg: dict, days: list[dict]) -> None:
    from docx import Document

    doc = Document(str(template))

    placeholders = {f"{{{{{k.upper()}}}}}": str(v) for k, v in cfg.items()}
    for p in _all_paragraphs(doc):
        for token, value in placeholders.items():
            _replace_in_paragraph(p, token, value)

    table = doc.tables[1]
    trs = table._tbl.findall(f"{W_NS}tr")
    proto = deepcopy(trs[-1])
    for tr in trs[1:]:
        table._tbl.remove(tr)

    proto_tcs = proto.findall(f"{W_NS}tc")
    proto_keg = proto_tcs[3].find(f"{W_NS}p")

    for day in days:
        tr = deepcopy(proto)
        tcs = tr.findall(f"{W_NS}tc")
        _set_cell_text(tcs[0], day["label"])
        _set_cell_text(tcs[1], day["in"])
        _set_cell_text(tcs[2], day["out"])

        tc = tcs[3]
        for p in tc.findall(f"{W_NS}p"):
            tc.remove(p)
        if day["items"]:
            for item in day["items"]:
                tc.append(_kegiatan_paragraph(proto_keg, item))
        else:
            empty = deepcopy(proto_keg)
            _set_left_alignment(empty)
            tc.append(empty)
        table._tbl.append(tr)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #

def find_serif_fonts() -> tuple[Path | None, Path | None]:
    pairs = [
        ("/usr/share/fonts/wps-fonts/times.ttf", "/usr/share/fonts/wps-fonts/timesbd.ttf"),
        ("/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf",
         "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"),
        ("/usr/share/fonts/liberation/LiberationSerif-Regular.ttf",
         "/usr/share/fonts/liberation/LiberationSerif-Bold.ttf"),
        ("/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
         "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf"),
        ("C:/Windows/Fonts/times.ttf", "C:/Windows/Fonts/timesbd.ttf"),
        ("/System/Library/Fonts/Supplemental/Times New Roman.ttf",
         "/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf"),
    ]
    for reg, bold in pairs:
        if Path(reg).exists() and Path(bold).exists():
            return Path(reg), Path(bold)
    return None, None


def render_pdf(out_path: Path, cfg: dict, days: list[dict], label: str,
               logo: Path | None) -> None:
    from fpdf import FPDF
    from fpdf.enums import TableBordersLayout, XPos, YPos
    from fpdf.fonts import FontFace

    reg, bold = find_serif_fonts()

    class LogBookPDF(FPDF):
        def __init__(self):
            super().__init__(orientation="P", unit="mm", format="A4")
            self.fam = "Times"
            if reg and bold:
                self.add_font("Book", "", str(reg))
                self.add_font("Book", "B", str(bold))
                self.fam = "Book"
            self.set_margins(25.4, 25.4, 25.4)
            self.set_auto_page_break(True, 22)

        def txt(self, s: str) -> str:
            if self.fam != "Times":
                return s
            return (s.replace("\u2014", "-").replace("\u2013", "-")
                     .replace("\u2018", "'").replace("\u2019", "'")
                     .replace("\u201c", '"').replace("\u201d", '"'))

        def header(self):
            if logo and Path(logo).exists():
                try:
                    self.image(str(logo), x=25.4, y=12.5, h=19)
                except Exception:
                    pass
            x0 = 55.0
            w = 210 - 25.4 - x0
            lines = (
                ("", 12, 5.4, "KEMENTERIAN PENDIDIKAN TINGGI, SAINS, DAN TEKNOLOGI"),
                ("", 11, 5.0, "POLITEKNIK NEGERI MALANG"),
                ("B", 11, 5.0, "JURUSAN TEKNOLOGI INFORMASI"),
                ("", 8.5, 4.0, "Jalan Soekarno Hatta Nomor 9, Jatimulyo, Lowokwaru, Malang 65141"),
                ("", 8.5, 4.0, "Telepon (0341) 404424, 404425, Faksimile (0341) 404420"),
                ("", 8.5, 4.0, "Laman www.polinema.ac.id"),
            )
            y0 = 12.5
            self.set_xy(x0, y0)
            for style, size, h, text in lines:
                self.set_x(x0)
                self.set_font(self.fam, style, size)
                self.multi_cell(w, h, text, align="C")
            y = max(self.get_y(), y0 + 19) + 1.2
            self.set_line_width(0.7)
            self.line(25.4, y, 210 - 25.4, y)
            self.set_y(y + 6.5)

    pdf = LogBookPDF()
    fam = pdf.fam
    pdf.add_page()

    pdf.set_font(fam, "B", 12)
    pdf.cell(0, 6, "LOG BOOK KEGIATAN", align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.cell(0, 6, "PROGRAM MAGANG INDUSTRI", align="C", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(4)

    pdf.set_font(fam, "B", 10.5)
    for key, label in (("nama", "Nama"), ("nim", "NIM"),
                       ("program_studi", "Program Studi"), ("mitra_industri", "Nama Mitra Industri")):
        pdf.cell(42, 5.4, label, new_x=XPos.RIGHT, new_y=YPos.TOP)
        pdf.cell(0, 5.4, pdf.txt(f": {cfg[key]}"), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(3.5)

    style_b = FontFace(emphasis="BOLD", family=fam, size_pt=10)
    pdf.set_font(fam, "", 10)
    with pdf.table(
        borders_layout=TableBordersLayout.ALL,
        col_widths=(31.8, 26.9, 24.5, 76.0),
        line_height=4.8,
        min_row_height=9,
        padding=1.2,
        first_row_as_headings=True,
        headings_style=style_b,
        text_align=("CENTER", "CENTER", "CENTER", "LEFT"),
    ) as table:
        head = table.row()
        for h in ("Hari, Tanggal", "Jam Masuk", "Jam Pulang", "Kegiatan"):
            head.cell(pdf.txt(h), align="CENTER")
        for day in days:
            row = table.row()
            row.cell(pdf.txt(day["label"]), style=style_b, align="CENTER")
            row.cell(day["in"], style=style_b, align="CENTER")
            row.cell(day["out"], style=style_b, align="CENTER")
            text = "\n".join(pdf.txt(f"- {i}") for i in day["items"]) if day["items"] else ""
            row.cell(text, align="LEFT")

    pdf.ln(7)
    if pdf.get_y() > 225:
        pdf.add_page()
    pdf.set_font(fam, "", 11)
    pdf.cell(0, 5.4, "Mahasiswa,", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(17)
    pdf.cell(0, 5.4, pdf.txt(cfg["nama"]), new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(9)
    pdf.cell(0, 5.4, "Mengetahui,", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    x = pdf.l_margin
    pdf.cell(80, 5.4, "Dosen Pembimbing,", new_x=XPos.RIGHT, new_y=YPos.TOP)
    pdf.set_x(x + 90)
    pdf.cell(0, 5.4, "Pembimbing Lapangan", new_x=XPos.LMARGIN, new_y=YPos.NEXT)
    pdf.ln(17)
    pdf.cell(80, 5.4, pdf.txt(cfg["dosen_pembimbing"]), new_x=XPos.RIGHT, new_y=YPos.TOP)
    pdf.set_x(x + 90)
    pdf.cell(0, 5.4, pdf.txt(cfg["pembimbing_lapangan"]), new_x=XPos.LMARGIN, new_y=YPos.NEXT)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    pdf.output(str(out_path))


# --------------------------------------------------------------------------- #
# main
# --------------------------------------------------------------------------- #

def load_config(path: Path, explicit: bool) -> dict:
    """Read the personal config. When the default workdir config is missing,
    create it from the shipped example and exit 2 so the user can edit it."""
    if not path.exists():
        if explicit:
            raise ValueError(
                f"config {path} tidak ditemukan (diberikan lewat --config)"
            )
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(EXAMPLE_CONFIG.read_text(encoding="utf-8"), encoding="utf-8")
        print(f"CONFIG_CREATED: {path} dibuat dari config.example.json.", file=sys.stderr)
        print("Edit datanya dulu (nama, nim, program_studi, mitra_industri, "
              "dosen_pembimbing, pembimbing_lapangan), lalu jalankan lagi.", file=sys.stderr)
        if (path.parent / ".git").exists():
            print(f"HINT: config ini berisi data pribadi — tambahkan "
                  f"'{path.name}' ke .gitignore di folder ini.", file=sys.stderr)
        sys.exit(2)

    try:
        cfg = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as e:
        raise ValueError(
            f"config {path} bukan JSON yang valid (baris {e.lineno}, kolom {e.colno}): {e.msg}"
        )
    if not isinstance(cfg, dict):
        raise ValueError(f"config {path} harus berupa objek JSON")

    required = ["nama", "nim", "program_studi", "mitra_industri",
                "dosen_pembimbing", "pembimbing_lapangan"]
    missing = [k for k in required if not str(cfg.get(k, "")).strip()]
    if missing:
        raise ValueError(f"config {path} kurang field: {', '.join(missing)}")
    return cfg


def main() -> int:
    ap = argparse.ArgumentParser(description="Build an internship log book (DOCX + PDF).")
    ap.add_argument("--input", default="", help='periode, mis. "September/minggu 2->3" atau "" (bulan berjalan)')
    ap.add_argument("--format", default="both", choices=["both", "word", "pdf"])
    ap.add_argument("--data", default=None, help="file JSON berisi todos & attendance dari MCP")
    ap.add_argument("--outdir", default=".", help="folder kerja (default: current dir)")
    ap.add_argument("--config", default=None,
                    help=f"path config (default: ./{CONFIG_FILENAME} di folder kerja)")
    ap.add_argument("--template", default=str(DEFAULT_TEMPLATE))
    ap.add_argument("--resolve-only", action="store_true",
                    help="hanya cetak resolusi periode (tanpa config/data/generate)")
    args = ap.parse_args()

    today = date.today()

    try:
        period = parse_period(args.input, today)
    except ValueError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        return 1

    folder_label, file_label = period_labels(period)

    if args.resolve_only:
        dates = []
        d = period["start"]
        while d <= period["end"]:
            dates.append(d.isoformat())
            d += timedelta(days=1)
        print(json.dumps({
            "ok": True,
            "mode": period["mode"],
            "start": period["start"].isoformat(),
            "end": period["end"].isoformat(),
            "label": file_label,
            "folder": f"Laporan {folder_label}",
            "weeks": list(period["weeks"]) if period["weeks"] else None,
            "dates": dates,
        }, ensure_ascii=False, indent=2))
        return 0

    explicit_config = args.config is not None
    config_path = Path(args.config).expanduser() if explicit_config \
        else Path.cwd() / CONFIG_FILENAME
    cfg = load_config(config_path, explicit_config)

    todos: list[dict] = []
    attendance: list[dict] = []
    if args.data:
        payload = json.loads(Path(args.data).read_text(encoding="utf-8"))
        todos = payload.get("todos") or []
        attendance = payload.get("attendance") or []

    days = build_days(period, todos, attendance)
    warnings = list(period["warnings"])

    outdir = Path(args.outdir).resolve()
    folder = outdir / f"Laporan {folder_label}"
    base = folder / f"Log Book {file_label}"

    files: list[str] = []
    overwritten = False

    if args.format in ("both", "word"):
        docx_path = base.with_suffix(".docx")
        overwritten = overwritten or docx_path.exists()
        fill_docx(Path(args.template), docx_path, cfg, days)
        files.append(str(docx_path))

    if args.format in ("both", "pdf"):
        pdf_path = base.with_suffix(".pdf")
        overwritten = overwritten or pdf_path.exists()
        try:
            render_pdf(pdf_path, cfg, days, file_label, DEFAULT_LOGO)
        except ImportError as e:
            print(f"ERROR: PDF butuh fpdf2 ({e}). Pakai --format word atau: pip install fpdf2",
                  file=sys.stderr)
            return 1
        files.append(str(pdf_path))

    no_hours = [d["date"].isoformat() for d in days if d["no_hours"]]
    no_activities = [d["date"].isoformat() for d in days if not d["items"]]

    print(json.dumps({
        "ok": True,
        "mode": period["mode"],
        "start": period["start"].isoformat(),
        "end": period["end"].isoformat(),
        "label": file_label,
        "folder": str(folder),
        "config": str(config_path),
        "files": files,
        "rows": len(days),
        "overwritten": overwritten,
        "days_without_hours": no_hours,
        "days_without_activities": no_activities,
        "warnings": warnings,
    }, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    try:
        sys.exit(main())
    except ValueError as e:
        print(f"ERROR: {e}", file=sys.stderr)
        sys.exit(1)
